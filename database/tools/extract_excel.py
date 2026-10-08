#!/usr/bin/env python3
"""
Ekstrak data mentah dari Excel "Laporan Simpin" ke JSON untuk skrip impor PHP.

  python database/tools/extract_excel.py "<file.xlsx>" database/import/excel-extract.json

Hanya MEMBACA Excel (nilai tersimpan, bukan perhitungan ulang). Selain data transaksi,
JSON memuat "checks": angka kontrol yang dihitung Excel sendiri. Impor PHP menolak commit
bila hasil di database tidak sama dengan angka kontrol ini.
"""
import datetime
import hashlib
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci

SHEET_SETORAN = "2-Laporan setoran "
SHEET_REKAP = "3-REKAP TABUNGAN+bunga"
SHEET_SHU = "6-JASA PINJAMAN & Kas SHU 20%"

# kolom awal blok bulanan di sheet setoran
MONTH_BLOCKS = {
    "2026-03": "B", "2026-04": "M", "2026-05": "W", "2026-06": "AG",
    "2026-07": "AS", "2026-08": "BC", "2026-09": "BN", "2026-10": "BZ",
    "2026-11": "CK", "2026-12": "CU", "2027-01": "DE", "2027-02": "DO",
}
FIRST_ROW, LAST_ROW, TOTAL_ROW = 9, 68, 69
CLOSING_ROW = 78  # "Saldo Akhir" ringkasan kas bulanan


def clean(v):
    return re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v


def num(v):
    return int(round(v)) if isinstance(v, (int, float)) and v else None


def iso(v):
    return v.date().isoformat() if isinstance(v, datetime.datetime) else None


def main(xlsx, out):
    # Excel di OneDrive kadang terkunci: salin dulu ke folder sementara.
    tmp = Path(tempfile.mkdtemp()) / "src.xlsx"
    shutil.copyfile(xlsx, tmp)
    sha = hashlib.sha256(tmp.read_bytes()).hexdigest()

    wv = openpyxl.load_workbook(tmp, data_only=True)
    wf = openpyxl.load_workbook(tmp)  # rumus, untuk mendeteksi sel yang diketik manual
    ws = wv[SHEET_SETORAN]

    # ---------- bulan ----------
    months = []
    for key, col in MONTH_BLOCKS.items():
        s = ci(col)
        meeting = None
        for c in range(s, s + 12):
            if isinstance(ws.cell(4, c).value, datetime.datetime):
                meeting = ws.cell(4, c).value
        location = None
        for c in range(s, s + 12):
            if clean(ws.cell(5, c).value) == "Lokasi":
                location = clean(ws.cell(5, c + 1).value) or None
        months.append({"key": key, "meeting_date": iso(meeting), "location": location})

    # ---------- anggota (sumber: blok Maret; baris 9..68) ----------
    members = []
    rekap = wf[SHEET_REKAP]
    for r in range(FIRST_ROW, LAST_ROW + 1):
        no, name = ws.cell(r, ci("B")).value, clean(ws.cell(r, ci("C")).value)
        if not name:
            continue
        aktif = clean(ws.cell(r, ci("F")).value) or "Maret 2026"
        bulan = {"Maret 2026": "2026-03", "April 2026": "2026-04"}.get(aktif)
        if bulan is None:
            raise SystemExit(f"Baris {r}: 'aktif per' tidak dikenal: {aktif!r}")
        # potongan cadangan 5% yang di-hardcode (bukan rumus) di sheet rekap tabungan
        bp = rekap.cell(r - 2, ci("BP")).value  # rekap baris 7 = anggota 1 = setoran baris 9
        exempt = bp is not None and not (isinstance(bp, str) and bp.startswith("="))
        members.append({
            "no": int(no), "name": name, "address": clean(ws.cell(r, ci("D")).value),
            "team": clean(ws.cell(r, ci("E")).value), "active_from": bulan + "-01",
            "reserve_exempt": bool(exempt), "excel_row": r,
        })

    team_names = []
    for m in members:
        if m["team"] not in team_names:
            team_names.append(m["team"])
    names = {m["name"]: m["no"] for m in members}
    teams = [{"name": t, "leader_no": names.get(t)} for t in team_names]

    # ---------- transaksi bulanan ----------
    monthly = {}
    check_month = {}
    closing = {}
    by_row = {m["excel_row"]: m["no"] for m in members}
    for key, col in MONTH_BLOCKS.items():
        s = ci(col)
        cols = {}
        for c in range(s, s + 12):
            h = ws.cell(7, c).value
            if h in ("MASUK", "KELUAR", "PINJAM", "TEMPO", "BAYAR"):
                cols[h] = c
        rows = []
        for r in range(FIRST_ROW, LAST_ROW + 1):
            if r not in by_row:
                continue
            item = {"no": by_row[r]}
            for h, field in (("MASUK", "tabungan"), ("KELUAR", "tarik"), ("PINJAM", "pinjam"), ("TEMPO", "tempo"), ("BAYAR", "bayar")):
                item[field] = num(ws.cell(r, cols[h]).value)
            if any(item[f] for f in ("tabungan", "tarik", "pinjam", "bayar")):
                rows.append(item)
        monthly[key] = rows
        check_month[key] = {
            "tabungan": num(ws.cell(TOTAL_ROW, cols["MASUK"]).value) or 0,
            "pinjam": num(ws.cell(TOTAL_ROW, cols["PINJAM"]).value) or 0,
            "bayar": num(ws.cell(TOTAL_ROW, cols["BAYAR"]).value) or 0,
        }
        # saldo akhir kas bulan itu: sel angka di baris 78 pada blok bulan
        vals = [ws.cell(CLOSING_ROW, c).value for c in range(s, s + 6)]
        vals = [v for v in vals if isinstance(v, (int, float))]
        closing[key] = int(round(vals[0])) if vals else None

    # ---------- biaya operasional (Kas SHU) ----------
    shu = wv[SHEET_SHU]
    expenses, current_month = [], None
    for r in range(5, 21):
        d, desc, amount = shu.cell(r, ci("N")).value, clean(shu.cell(r, ci("O")).value), shu.cell(r, ci("Q")).value
        if isinstance(d, datetime.datetime):
            current_month = d.strftime("%Y-%m")
        if isinstance(amount, (int, float)) and amount > 0 and desc:
            expenses.append({"month": current_month, "description": desc, "amount": int(round(amount))})

    # ---------- angka kontrol dari Excel ----------
    sep = ws  # sheet setoran
    member_savings = {}
    member_outstanding = {}
    for m in members:
        r = m["excel_row"]
        # total tabungan per anggota: rekap tabungan kolom BN ("TABUNGAN" akhir)
        v = wv[SHEET_REKAP].cell(r - 2, ci("BN")).value
        member_savings[str(m["no"])] = int(round(v)) if isinstance(v, (int, float)) else 0
        # sisa pinjaman per anggota: rekap pinjaman kolom FX di sheet setoran
        o = sep.cell(r, ci("FX")).value
        member_outstanding[str(m["no"])] = int(round(o)) if isinstance(o, (int, float)) else 0

    checks = {
        "month_totals": check_month,
        "cash_closing": closing,
        "member_savings": member_savings,
        "member_outstanding": member_outstanding,
        "total_pinjaman": num(wv[SHEET_SHU]["F19"].value),
        "total_bunga": num(wv[SHEET_SHU]["G19"].value),
        "total_biaya": num(wv[SHEET_SHU]["Q21"].value),
    }

    data = {
        "source": {"file": Path(xlsx).name, "sha256": sha, "extracted_at": datetime.datetime.now().isoformat(timespec="seconds")},
        "period": {"name": "THR Adem Ayem 2026-2027", "start": "2026-03-01", "end": "2027-02-28"},
        "rate_pct_month": 2,
        "months": months, "teams": teams, "members": members,
        "monthly": monthly, "expenses": expenses, "checks": checks,
    }
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    Path(out).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"OK  {len(members)} anggota, {len(teams)} regu, "
          f"{sum(len(v) for v in monthly.values())} baris bulanan, {len(expenses)} biaya -> {out}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
